"""Multi-format text and structural metadata extractor for PDF, DOCX, PPTX, TXT, and MD."""
import os
from typing import List, Dict, Any


class ExtractedSection:
    def __init__(self, text: str, page_number: int | None = None, section_title: str | None = None):
        self.text = text.strip()
        self.page_number = page_number
        self.section_title = section_title

    def to_dict(self) -> Dict[str, Any]:
        return {
            "text": self.text,
            "page_number": self.page_number,
            "section_title": self.section_title,
        }


class TextExtractor:
    @classmethod
    def extract_from_file(cls, file_path: str) -> List[ExtractedSection]:
        """Extracts text sections with page/section metadata based on file extension."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return cls._extract_pdf(file_path)
        elif ext == ".docx":
            return cls._extract_docx(file_path)
        elif ext == ".pptx":
            return cls._extract_pptx(file_path)
        elif ext in (".txt", ".md"):
            return cls._extract_text_or_markdown(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Supported: .pdf, .docx, .pptx, .txt, .md")

    @classmethod
    def _extract_pdf(cls, file_path: str) -> List[ExtractedSection]:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        sections = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                sections.append(ExtractedSection(
                    text=text,
                    page_number=i + 1,
                    section_title=f"Page {i + 1}"
                ))
        return sections

    @classmethod
    def _extract_docx(cls, file_path: str) -> List[ExtractedSection]:
        import docx
        doc = docx.Document(file_path)
        sections = []
        current_heading = "Introduction"
        current_paragraphs: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style and p.style.name.startswith("Heading"):
                if current_paragraphs:
                    sections.append(ExtractedSection(
                        text="\n\n".join(current_paragraphs),
                        section_title=current_heading
                    ))
                    current_paragraphs = []
                current_heading = text
            else:
                current_paragraphs.append(text)

        if current_paragraphs:
            sections.append(ExtractedSection(
                text="\n\n".join(current_paragraphs),
                section_title=current_heading
            ))
        return sections

    @classmethod
    def _extract_pptx(cls, file_path: str) -> List[ExtractedSection]:
        import pptx
        prs = pptx.Presentation(file_path)
        sections = []
        for i, slide in enumerate(prs.slides):
            slide_texts: List[str] = []
            slide_title = f"Slide {i + 1}"
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    if shape == slide.shapes[0] and len(shape.text.strip()) < 80:
                        slide_title = shape.text.strip()
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                sections.append(ExtractedSection(
                    text="\n".join(slide_texts),
                    page_number=i + 1,
                    section_title=slide_title
                ))
        return sections

    @classmethod
    def _extract_text_or_markdown(cls, file_path: str) -> List[ExtractedSection]:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        sections = []
        lines = content.split("\n")
        current_heading = os.path.splitext(os.path.basename(file_path))[0]
        current_lines: List[str] = []
        page_approx = 1

        for line in lines:
            stripped = line.strip()
            if stripped.startswith("#"):
                if current_lines:
                    text = "\n".join(current_lines).strip()
                    if text:
                        sections.append(ExtractedSection(
                            text=text,
                            page_number=page_approx,
                            section_title=current_heading
                        ))
                        page_approx += 1
                    current_lines = []
                current_heading = stripped.lstrip("#").strip()
            else:
                current_lines.append(line)

        if current_lines:
            text = "\n".join(current_lines).strip()
            if text:
                sections.append(ExtractedSection(
                    text=text,
                    page_number=page_approx,
                    section_title=current_heading
                ))

        return sections
